#!/usr/bin/env python3
"""Convert the handbook's download artifacts (public/artifacts/*.md) to .docx.

Handles the markdown actually used in these files: # / ## / ### headings,
**bold**, `inline code`, `- `/`* ` bullets, `1.` numbered lists,
`- [ ]`/`- [x]` checkboxes (rendered as readable ☐/☑), pipe tables, `---`
horizontal rules, and plain paragraphs. HTML comments are stripped.

Usage: python scripts/md_to_docx.py [artifacts_dir]
# ponytail: line-based parser, not a full CommonMark impl — covers the
# patterns these 41 files use; anything fancier degrades to plain text.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

CODE_FONT = "Consolas"
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
HTML_COMMENT = re.compile(r"<!--.*?-->", re.DOTALL)
BULLET = re.compile(r"^(\s*)[-*]\s+(.*)$")
NUMBERED = re.compile(r"^(\s*)\d+\.\s+(.*)$")
CHECKBOX = re.compile(r"^(\s*)[-*]\s+\[([ xX])\]\s+(.*)$")
HEADING = re.compile(r"^(#{1,6})\s+(.*)$")


def add_runs(paragraph, text):
    """Add text to a paragraph, honoring **bold** and `code` spans."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = CODE_FONT
        else:
            paragraph.add_run(part)


def flush_table(doc, rows):
    """rows: list of lists of cell strings (separator row already removed)."""
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.paragraphs[0].text = ""
            add_runs(cell.paragraphs[0], row[j] if j < len(row) else "")
            if i == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def split_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def is_sep_row(cells):
    return cells and all(re.fullmatch(r":?-+:?", c) for c in cells if c != "")


def convert(md_path: Path, docx_path: Path):
    text = HTML_COMMENT.sub("", md_path.read_text(encoding="utf-8"))
    lines = text.split("\n")
    doc = Document()
    table_buf = []

    def flush():
        if table_buf:
            rows = [r for r in table_buf if not is_sep_row(r)]
            flush_table(doc, rows)
            table_buf.clear()

    for line in lines:
        stripped = line.strip()

        # table lines accumulate until a non-table line
        if stripped.startswith("|"):
            table_buf.append(split_row(line))
            continue
        flush()

        if not stripped:
            continue

        if stripped == "---" or re.fullmatch(r"[-*_]{3,}", stripped):
            doc.add_paragraph().add_run("―" * 20)  # horizontal bar
            continue

        m = HEADING.match(stripped)
        if m:
            level = min(len(m.group(1)), 3)
            doc.add_heading(m.group(2), level=level)
            continue

        m = CHECKBOX.match(line)
        if m:
            box = "☑" if m.group(2).lower() == "x" else "☐"
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(box + " ")
            add_runs(p, m.group(3))
            continue

        m = NUMBERED.match(line)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, m.group(2))
            continue

        m = BULLET.match(line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, m.group(2))
            continue

        p = doc.add_paragraph()
        add_runs(p, stripped)

    flush()
    doc.save(str(docx_path))


def main():
    base = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("public/artifacts")
    md_files = sorted(base.glob("*.md"))
    for md in md_files:
        convert(md, md.with_suffix(".docx"))
    print(f"converted {len(md_files)} files")


if __name__ == "__main__":
    main()
